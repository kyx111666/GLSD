from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
slide=prs.slides.add_slide(prs.slide_layouts[6])
FONT='Times New Roman'
COL={'navy':RGBColor(35,63,91),'blue':RGBColor(72,125,170),'bluefill':RGBColor(232,242,249),'green':RGBColor(79,137,108),'greenfill':RGBColor(233,245,237),'purple':RGBColor(112,96,150),'gray':RGBColor(90,98,108),'line':RGBColor(90,105,120),'bg':RGBColor(249,250,252),'gold':RGBColor(160,125,55)}

def rect(x,y,w,h,fill,line=COL['line'],radius=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line; sh.line.width=Pt(1.2); return sh

def txt(x,y,w,h,s,size=12,color=COL['navy'],bold=False,align=PP_ALIGN.CENTER):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=s; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return tb

def arrow(x1,y1,x2,y2,color=COL['line'],width=1.5):
    c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2)); c.line.color.rgb=color; c.line.width=Pt(width); c.line.end_arrowhead=True; return c

def wave(x,y,w,h,color):
    # editable polyline-like segments
    pts=[(0,.65),(.12,.45),(.24,.72),(.38,.3),(.52,.6),(.68,.2),(.82,.5),(1,.35)]
    for (a,b),(c,d) in zip(pts,pts[1:]):
        ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+a*w),Inches(y+b*h),Inches(x+c*w),Inches(y+d*h)); ln.line.color.rgb=color; ln.line.width=Pt(1.2)

# background and title
rect(0,0,13.333,7.5,COL['bg'],COL['bg'])
txt(.35,.18,12.6,.38,'GLSD (Global–Local Saliency Decoder): Method Overview',22,COL['navy'],True,PP_ALIGN.LEFT)
# major zones
rect(.3,.75,2.35,5.35,RGBColor(246,248,251),COL['line'])
rect(2.85,.75,7.85,5.35,RGBColor(248,250,252),COL['line'])
rect(10.95,.75,2.08,5.35,RGBColor(246,248,251),COL['line'])
txt(.45,.82,2.1,.3,'Frozen backbone inputs',13,COL['navy'],True)
txt(3.0,.82,7.5,.3,'Shared GLSD core',13,COL['navy'],True)
txt(11.1,.82,1.8,.3,'Decoded outputs',13,COL['navy'],True)
# inputs
for y,label,col in [(1.35,'ME-TST+ cached 1D\nspotting response',COL['blue']),(2.5,'BoostingVRME cached\ntemporal curve',COL['green'])]:
    rect(.5,y,1.85,.72,RGBColor(255,255,255),col,True); txt(.58,y+.08,1.68,.5,label,10,COL['navy'],True); wave(.68,y+.55,1.45,.18,col)
    rect(1.0,y+.88,.85,.42,RGBColor(255,255,255),COL['line'],True); txt(1.02,y+.9,.81,.35,'adapter',10,COL['gray'])
    arrow(1.43,y+.72,1.43,y+.88,col)
arrow(1.85,2.98,2.65,2.98,COL['line']); txt(1.45,3.15,1.25,.48,'unified frozen\nresponse x(t)',10,COL['navy'],True)
# core flow boxes
boxes=[('Multi-scale smoothing\ny1.0(t), y1.5(t), y2.0(t)',3.15,1.35,1.42),('Reference candidate peaks\nat scale a0',4.75,1.35,1.42),('Thresholding\nCτ={c:S(c)≥τ}',8.75,1.35,1.42),('Event selection / decoding',8.75,3.05,1.42)]
for label,x,y,w in boxes: rect(x,y,w,.85,RGBColor(255,255,255),COL['purple'],True); txt(x+.05,y+.1,w-.1,.65,label,10,COL['navy'],True)
arrow(2.65,2.98,3.15,1.78); arrow(4.57,1.78,4.75,1.78); arrow(6.17,1.78,6.55,1.78)
# branches
rect(6.55,1.05,1.85,1.45,COL['bluefill'],COL['blue'],True); txt(6.68,1.15,1.6,.35,'Global prominence',12,COL['blue'],True); txt(6.68,1.55,1.6,.45,'evidence  G(c)',14,COL['blue'],True)
rect(6.55,2.8,1.85,1.45,COL['greenfill'],COL['green'],True); txt(6.68,2.9,1.6,.35,'Aligned local',12,COL['green'],True); txt(6.68,3.3,1.6,.45,'multi-scale saliency  L(c)',12,COL['green'],True)
arrow(6.17,1.78,6.55,1.78,COL['blue']); arrow(6.17,1.78,6.55,3.5,COL['green'])
rect(6.25,4.65,2.55,.9,RGBColor(255,250,235),COL['gold'],True); txt(6.35,4.78,2.35,.55,'Composite score\nS(c)=G(c)+L(c)²',12,COL['gold'],True)
arrow(7.48,2.5,7.48,4.65,COL['blue']); arrow(7.48,4.25,7.48,4.65,COL['green']); arrow(8.8,5.1,9.45,3.9,COL['gold']); arrow(9.45,2.2,8.75,1.78,COL['purple']); arrow(9.46,3.9,9.45,3.9,COL['purple'])
# outputs
for y,label,col in [(1.55,'ME-TST+\nevent intervals',COL['blue']),(3.0,'BoostingVRME\nevent intervals',COL['green'])]:
    rect(11.15,y,1.68,.85,RGBColor(255,255,255),col,True); txt(11.22,y+.12,1.54,.58,label,10,COL['navy'],True)
arrow(10.17,1.78,11.15,1.98,COL['blue']); arrow(10.17,3.48,11.15,3.4,COL['green']); txt(11.05,4.25,1.85,.6,'Output rule set by\ncorresponding adapter',9,COL['gray'])
# bottom notes
rect(.45,6.28,12.4,.82,RGBColor(241,244,247),COL['line'])
txt(.65,6.37,5.9,.55,'Shared GLSD core: candidate construction, G/L computation, cross-scale alignment, composite scoring, candidate filtering',10,COL['navy'],False,PP_ALIGN.LEFT)
txt(6.7,6.37,5.9,.55,'Method-specific adapters provide response format, base time scale k, event geometry, and conflict rules',10,COL['navy'],False,PP_ALIGN.LEFT)
# small labels
for x,y,s,c in [(6.62,2.55,'global',COL['blue']),(6.62,4.3,'local',COL['green'])]: txt(x,y,.55,.2,s,8,c,True,PP_ALIGN.LEFT)
prs.save('GLSD_method_overview_editable.pptx')
